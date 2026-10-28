from __future__ import annotations

# AgentChat/direct_codex 2026-07-13T01:28:53-04:00: bound Slack file ingest and validate downloaded content signatures.

import hashlib
import io
import os
import re
import urllib.request
import zipfile
from pathlib import Path
from typing import Any

from .config import BridgeConfig
from .memory import redact_secrets


SAFE_NAME_RE = re.compile(r"[^A-Za-z0-9._ -]+")
TEXT_EXTENSIONS = {
    ".cfg",
    ".conf",
    ".css",
    ".csv",
    ".htm",
    ".html",
    ".ini",
    ".js",
    ".json",
    ".jsonl",
    ".log",
    ".md",
    ".ps1",
    ".py",
    ".srt",
    ".sql",
    ".ts",
    ".tsv",
    ".txt",
    ".vtt",
    ".xml",
    ".yaml",
    ".yml",
}
TEXT_MIMETYPE_PREFIXES = ("text/",)
TEXT_MIMETYPES = {
    "application/csv",
    "application/json",
    "application/sql",
    "application/xml",
    "application/x-ndjson",
    "application/x-powershell",
    "application/x-yaml",
}
IMAGE_EXTENSIONS = {".bmp", ".gif", ".heic", ".heif", ".jpeg", ".jpg", ".png", ".tif", ".tiff", ".webp"}
PDF_EXTENSIONS = {".pdf"}
DOCX_EXTENSIONS = {".docx"}
PPTX_EXTENSIONS = {".pptx"}
SPREADSHEET_EXTENSIONS = {".xlsx", ".xlsm"}
LEGACY_SPREADSHEET_EXTENSIONS = {".xls"}
ZIP_EXTENSIONS = {".zip"}
RETURN_FILE_EXTENSIONS = frozenset(
    {
        ".bmp",
        ".csv",
        ".docx",
        ".gif",
        ".heic",
        ".heif",
        ".jpeg",
        ".jpg",
        ".json",
        ".md",
        ".pdf",
        ".png",
        ".pptx",
        ".tif",
        ".tiff",
        ".tsv",
        ".txt",
        ".webp",
        ".xls",
        ".xlsm",
        ".xlsx",
    }
)
UPLOAD_DIRECTIVE_RE = re.compile(
    r"(?im)^\s*BRIDGE_UPLOAD(?:_FILE)?:\s*(?P<path>.+?)\s*$"
)


def download_event_files(
    *,
    config: BridgeConfig,
    event: dict[str, Any],
    client: Any,
) -> list[dict[str, Any]]:
    if not config.defaults.get("slack_file_downloads_enabled", True):
        return []

    files = _event_files(event)
    max_files = max(1, _config_int(config.defaults, "slack_file_max_files_per_event", 5))
    files = files[:max_files]
    if not files:
        return []

    root = Path(str(config.defaults.get("slack_file_download_dir", "data/slack_files")))
    channel = safe_filename(str(event.get("channel") or "unknown-channel"))
    ts = safe_filename(str(event.get("ts") or "unknown-ts").replace(".", "-"))
    target_dir = root / channel / ts
    target_dir.mkdir(parents=True, exist_ok=True)

    token = getattr(client, "token", None) or os.getenv("SLACK_BOT_TOKEN")
    max_bytes = int(float(config.defaults.get("slack_file_max_download_mb", 25)) * 1024 * 1024)
    preview_chars = _config_int(config.defaults, "slack_file_text_preview_chars", 6000)
    pdf_max_pages = _config_int(config.defaults, "slack_file_pdf_max_pages", 10)
    spreadsheet_max_rows = _config_int(config.defaults, "slack_file_spreadsheet_preview_rows", 80)
    spreadsheet_max_cols = _config_int(config.defaults, "slack_file_spreadsheet_preview_cols", 30)
    download_timeout = _config_int(config.defaults, "slack_file_download_timeout_seconds", 20)
    ocr_timeout = _config_int(config.defaults, "slack_file_ocr_timeout_seconds", 15)
    ocr_enabled = _config_bool(config.defaults, "slack_file_ocr_enabled", True)
    pdf_ocr_enabled = _config_bool(config.defaults, "slack_file_pdf_ocr_enabled", True)
    downloaded: list[dict[str, Any]] = []

    for raw_file_meta in files:
        file_meta = _hydrate_file_meta(raw_file_meta, client)
        record = _file_record(file_meta)
        filename = safe_filename(record["name"] or record["id"] or "slack-file")
        destination = unique_path(target_dir / filename)
        url = file_meta.get("url_private_download") or file_meta.get("url_private")

        try:
            if url and token:
                _download_private_url(
                    url,
                    destination,
                    token,
                    max_bytes,
                    timeout_seconds=download_timeout,
                    expected_mimetype=str(record.get("mimetype") or ""),
                )
                record["local_path"] = str(destination.resolve())
                record["sha256"] = file_sha256(destination)
                record["preview"] = preview_file_text(
                    destination,
                    record,
                    preview_chars,
                    ocr_enabled=ocr_enabled,
                    pdf_ocr_enabled=pdf_ocr_enabled,
                    pdf_max_pages=pdf_max_pages,
                    spreadsheet_max_rows=spreadsheet_max_rows,
                    spreadsheet_max_cols=spreadsheet_max_cols,
                    ocr_timeout=ocr_timeout,
                )
            elif file_meta.get("preview"):
                record["preview"] = redact_secrets(str(file_meta.get("preview"))[:preview_chars])
                record["download_error"] = "Slack file URL or bot token was unavailable."
            else:
                record["download_error"] = "Slack file could not be downloaded."
        except Exception as exc:
            record["download_error"] = str(exc)
        downloaded.append(record)

    return downloaded


def format_files_for_prompt(files: list[dict[str, Any]] | None) -> str:
    if not files:
        return ""

    lines = ["Attached Slack files:"]
    for index, item in enumerate(files, start=1):
        name = item.get("name") or item.get("title") or item.get("id") or f"file-{index}"
        local_path = item.get("local_path")
        mimetype = item.get("mimetype") or "unknown type"
        size = item.get("size")
        lines.append(f"- {name} ({mimetype}, {size or 'unknown'} bytes)")
        if local_path:
            lines.append(f"  Local path: {local_path}")
        if item.get("sha256"):
            lines.append(f"  SHA256: {item['sha256']}")
        if item.get("download_error"):
            lines.append(f"  Download issue: {item['download_error']}")
        for note in item.get("preview_notes") or []:
            lines.append(f"  Decode note: {note}")
        preview = str(item.get("preview") or "").strip()
        if preview:
            lines.append("  Decoded preview:")
            for line in preview.splitlines()[:80]:
                lines.append(f"    {line}")

    return "\n".join(lines).strip() + "\n\n"


def event_has_files(event: dict[str, Any]) -> bool:
    return bool(_event_files(event))


def split_upload_directives(text: str) -> tuple[str, list[str]]:
    paths: list[str] = []

    def replace(match: re.Match[str]) -> str:
        raw = match.group("path").strip().strip("\"'")
        if raw:
            paths.append(raw)
        return ""

    cleaned = UPLOAD_DIRECTIVE_RE.sub(replace, text or "")
    cleaned = re.sub(r"\n{3,}", "\n\n", cleaned).strip()
    return cleaned, paths


def safe_filename(value: str) -> str:
    clean = SAFE_NAME_RE.sub("_", value).strip(" .")
    return clean[:160] or "file"


def unique_path(path: Path) -> Path:
    if not path.exists():
        return path
    stem = path.stem
    suffix = path.suffix
    parent = path.parent
    for index in range(1, 1000):
        candidate = parent / f"{stem}-{index}{suffix}"
        if not candidate.exists():
            return candidate
    raise FileExistsError(f"Could not find an unused path for {path}")


def preview_file_text(
    path: Path,
    record: dict[str, Any],
    max_chars: int,
    *,
    ocr_enabled: bool = True,
    pdf_ocr_enabled: bool = True,
    pdf_max_pages: int = 10,
    spreadsheet_max_rows: int = 80,
    spreadsheet_max_cols: int = 30,
    ocr_timeout: int = 15,
) -> str:
    if max_chars <= 0:
        return ""

    preview, notes = _decode_file_preview(
        path,
        record,
        max_chars=max_chars,
        ocr_enabled=ocr_enabled,
        pdf_ocr_enabled=pdf_ocr_enabled,
        pdf_max_pages=pdf_max_pages,
        spreadsheet_max_rows=spreadsheet_max_rows,
        spreadsheet_max_cols=spreadsheet_max_cols,
        ocr_timeout=ocr_timeout,
    )
    if notes:
        record["preview_notes"] = notes
    return redact_secrets(_limit_text(preview, max_chars))


def _decode_file_preview(
    path: Path,
    record: dict[str, Any],
    *,
    max_chars: int,
    ocr_enabled: bool,
    pdf_ocr_enabled: bool,
    pdf_max_pages: int,
    spreadsheet_max_rows: int,
    spreadsheet_max_cols: int,
    ocr_timeout: int,
) -> tuple[str, list[str]]:
    suffix = path.suffix.lower()
    mimetype = str(record.get("mimetype") or "").lower()

    if _is_text_file(suffix, mimetype):
        return _read_text_preview(path, max_chars), []
    if suffix in PDF_EXTENSIONS or mimetype == "application/pdf":
        return _extract_pdf_text(
            path,
            max_chars=max_chars,
            max_pages=pdf_max_pages,
            ocr_enabled=ocr_enabled and pdf_ocr_enabled,
            ocr_timeout=ocr_timeout,
        )
    if suffix in DOCX_EXTENSIONS or mimetype.endswith("wordprocessingml.document"):
        return _extract_docx_text(path, max_chars=max_chars)
    if suffix in PPTX_EXTENSIONS or mimetype.endswith("presentationml.presentation"):
        return _extract_pptx_text(path, max_chars=max_chars)
    if suffix in SPREADSHEET_EXTENSIONS or mimetype.endswith("spreadsheetml.sheet"):
        return _extract_spreadsheet_text(
            path,
            max_chars=max_chars,
            max_rows=spreadsheet_max_rows,
            max_cols=spreadsheet_max_cols,
        )
    if suffix in IMAGE_EXTENSIONS or mimetype.startswith("image/"):
        if not ocr_enabled:
            return "", ["Image OCR is disabled by bridge config."]
        return _ocr_image_file(path, max_chars=max_chars, ocr_timeout=ocr_timeout)
    if suffix in ZIP_EXTENSIONS or mimetype in {"application/zip", "application/x-zip-compressed"}:
        return _list_zip_contents(path, max_chars=max_chars)
    if suffix == ".xls":
        return "", ["Legacy .xls files are saved locally but not decoded. Upload CSV or XLSX for preview text."]

    label = mimetype or suffix or "unknown type"
    return "", [f"No text decoder is configured for {label}. The local file path is available above."]


def _is_text_file(suffix: str, mimetype: str) -> bool:
    return (
        suffix in TEXT_EXTENSIONS
        or mimetype in TEXT_MIMETYPES
        or mimetype.startswith(TEXT_MIMETYPE_PREFIXES)
    )


def _read_text_preview(path: Path, max_chars: int) -> str:
    raw = path.read_bytes()[: max(max_chars * 4, 8192)]
    for encoding in ("utf-8-sig", "utf-16", "utf-16-le", "utf-16-be", "cp1252"):
        try:
            return raw.decode(encoding, errors="strict")
        except UnicodeError:
            continue
    return raw.decode("utf-8", errors="replace")


def _extract_pdf_text(
    path: Path,
    *,
    max_chars: int,
    max_pages: int,
    ocr_enabled: bool,
    ocr_timeout: int,
) -> tuple[str, list[str]]:
    notes: list[str] = []
    try:
        from pypdf import PdfReader
    except Exception:
        return "", ["PDF text extraction needs the pypdf package. Run scripts/install.ps1 after updating dependencies."]

    try:
        reader = PdfReader(str(path))
    except Exception as exc:
        return "", [f"PDF could not be opened: {exc}"]

    page_count = len(reader.pages)
    page_limit = max(1, min(max_pages, page_count))
    parts: list[str] = []
    for index in range(page_limit):
        try:
            page_text = reader.pages[index].extract_text() or ""
        except Exception as exc:
            notes.append(f"PDF page {index + 1} text extraction failed: {exc}")
            continue
        page_text = page_text.strip()
        if page_text:
            parts.append(f"[Page {index + 1}]\n{page_text}")
        if len("\n\n".join(parts)) >= max_chars:
            break

    if page_count > page_limit:
        notes.append(f"PDF preview limited to first {page_limit} of {page_count} pages.")
    text = "\n\n".join(parts).strip()
    if text:
        return text, notes

    notes.append("PDF has no selectable text in the previewed pages.")
    if not ocr_enabled:
        return "", notes

    ocr_text, ocr_notes = _ocr_pdf_pages(
        path,
        max_chars=max_chars,
        max_pages=page_limit,
        ocr_timeout=ocr_timeout,
    )
    return ocr_text, notes + ocr_notes


def _ocr_pdf_pages(
    path: Path,
    *,
    max_chars: int,
    max_pages: int,
    ocr_timeout: int,
) -> tuple[str, list[str]]:
    notes: list[str] = []
    try:
        import fitz
    except Exception:
        return "", ["Scanned PDF OCR needs PyMuPDF plus Tesseract OCR."]

    try:
        document = fitz.open(str(path))
    except Exception as exc:
        return "", [f"Scanned PDF OCR could not open the file: {exc}"]

    parts: list[str] = []
    try:
        for index in range(min(max_pages, len(document))):
            try:
                page = document[index]
                pixmap = page.get_pixmap(matrix=fitz.Matrix(2, 2), alpha=False)
                image = _image_from_bytes(pixmap.tobytes("png"))
                remaining_chars = max_chars - len("\n\n".join(parts))
                if remaining_chars <= 0:
                    break
                page_text, page_notes = _ocr_image(
                    image,
                    max_chars=remaining_chars,
                    ocr_timeout=ocr_timeout,
                )
                notes.extend(f"PDF page {index + 1}: {note}" for note in page_notes)
                if page_text.strip():
                    parts.append(f"[OCR page {index + 1}]\n{page_text.strip()}")
                if len("\n\n".join(parts)) >= max_chars:
                    break
            except Exception as exc:
                notes.append(f"PDF page {index + 1} OCR failed: {exc}")
    finally:
        document.close()

    return "\n\n".join(parts).strip(), notes


def _extract_docx_text(path: Path, *, max_chars: int) -> tuple[str, list[str]]:
    try:
        from docx import Document
    except Exception:
        return "", ["DOCX extraction needs the python-docx package. Run scripts/install.ps1 after updating dependencies."]

    try:
        document = Document(str(path))
    except Exception as exc:
        return "", [f"DOCX could not be opened: {exc}"]

    parts: list[str] = []
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if text:
            parts.append(text)
        if len("\n".join(parts)) >= max_chars:
            break
    for table_index, table in enumerate(document.tables, start=1):
        if len("\n".join(parts)) >= max_chars:
            break
        parts.append(f"[Table {table_index}]")
        for row in table.rows:
            values = [cell.text.strip().replace("\n", " ") for cell in row.cells]
            if any(values):
                parts.append(" | ".join(values))
            if len("\n".join(parts)) >= max_chars:
                break
    return "\n".join(parts).strip(), []


def _extract_pptx_text(path: Path, *, max_chars: int) -> tuple[str, list[str]]:
    try:
        from pptx import Presentation
    except Exception:
        return "", ["PPTX extraction needs the python-pptx package. Run scripts/install.ps1 after updating dependencies."]

    try:
        presentation = Presentation(str(path))
    except Exception as exc:
        return "", [f"PPTX could not be opened: {exc}"]

    parts: list[str] = []
    for slide_index, slide in enumerate(presentation.slides, start=1):
        slide_lines: list[str] = []
        for shape in slide.shapes:
            text = getattr(shape, "text", "")
            if text and text.strip():
                slide_lines.append(text.strip())
        if slide_lines:
            parts.append(f"[Slide {slide_index}]\n" + "\n".join(slide_lines))
        if len("\n\n".join(parts)) >= max_chars:
            break
    return "\n\n".join(parts).strip(), []


def _extract_spreadsheet_text(
    path: Path,
    *,
    max_chars: int,
    max_rows: int,
    max_cols: int,
) -> tuple[str, list[str]]:
    notes: list[str] = []
    try:
        from openpyxl import load_workbook
    except Exception:
        return "", ["XLSX extraction needs the openpyxl package. Run scripts/install.ps1 after updating dependencies."]

    try:
        workbook = load_workbook(filename=str(path), read_only=True, data_only=True)
    except Exception as exc:
        return "", [f"Spreadsheet could not be opened: {exc}"]

    parts: list[str] = []
    rows_seen = 0
    try:
        for worksheet in workbook.worksheets:
            parts.append(f"[Sheet: {worksheet.title}]")
            for row in worksheet.iter_rows(max_col=max_cols, values_only=True):
                values = ["" if value is None else str(value).replace("\n", " ") for value in row]
                if not any(value.strip() for value in values):
                    continue
                parts.append("\t".join(values).rstrip())
                rows_seen += 1
                if rows_seen >= max_rows or len("\n".join(parts)) >= max_chars:
                    notes.append(f"Spreadsheet preview limited to {rows_seen} non-empty rows.")
                    return "\n".join(parts).strip(), notes
    finally:
        workbook.close()

    return "\n".join(parts).strip(), notes


def _ocr_image_file(path: Path, *, max_chars: int, ocr_timeout: int) -> tuple[str, list[str]]:
    try:
        from PIL import Image, ImageOps
    except Exception:
        return "", ["Image OCR needs Pillow plus Tesseract OCR."]

    try:
        with Image.open(path) as image:
            image = ImageOps.exif_transpose(image)
            image.load()
            return _ocr_image(image, max_chars=max_chars, ocr_timeout=ocr_timeout)
    except Exception as exc:
        return "", [f"Image could not be opened for OCR: {exc}"]


def _image_from_bytes(data: bytes) -> Any:
    from PIL import Image

    image = Image.open(io.BytesIO(data))
    image.load()
    return image


def _ocr_image(image: Any, *, max_chars: int, ocr_timeout: int) -> tuple[str, list[str]]:
    notes: list[str] = []
    try:
        import pytesseract
    except Exception:
        return "", ["Image OCR needs pytesseract and the Tesseract OCR executable."]

    max_pixels = 12_000_000
    width, height = image.size
    if width * height > max_pixels:
        image = image.copy()
        image.thumbnail((4000, 4000))
        notes.append("Image was downscaled before OCR.")

    try:
        text = pytesseract.image_to_string(image, timeout=max(1, ocr_timeout)) or ""
    except Exception as exc:
        return "", [f"Image OCR failed: {exc}"]

    text = text.strip()
    if not text:
        notes.append("Image OCR found no text.")
    return _limit_text(text, max_chars), notes


def _list_zip_contents(path: Path, *, max_chars: int) -> tuple[str, list[str]]:
    try:
        with zipfile.ZipFile(path) as archive:
            names = archive.namelist()
    except Exception as exc:
        return "", [f"ZIP could not be opened: {exc}"]

    preview = "\n".join(names[:200])
    notes: list[str] = []
    if len(names) > 200:
        notes.append(f"ZIP listing limited to 200 of {len(names)} entries.")
    return _limit_text(preview, max_chars), notes


def _limit_text(text: str, max_chars: int) -> str:
    if len(text) <= max_chars:
        return text
    return text[:max_chars].rstrip() + "\n[preview truncated]"


def file_sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _download_private_url(
    url: str,
    destination: Path,
    token: str,
    max_bytes: int,
    *,
    timeout_seconds: int,
    expected_mimetype: str = "",
) -> None:
    request = urllib.request.Request(url, headers={"Authorization": f"Bearer {token}"})
    try:
        with urllib.request.urlopen(request, timeout=max(1, timeout_seconds)) as response:
            content_type = str(response.headers.get("Content-Type") or "").split(";", 1)[0].strip().lower()
            expected_type = str(expected_mimetype or "").split(";", 1)[0].strip().lower()
            expected_is_html = expected_type in {"text/html", "application/xhtml+xml"}
            destination_is_html = destination.suffix.lower() in {".htm", ".html"}
            if content_type in {"text/html", "application/xhtml+xml"} and not (
                expected_is_html or destination_is_html
            ):
                raise ValueError(
                    "Slack returned an HTML sign-in page instead of the requested file. "
                    "The installed Slack app must include files:read and be reinstalled to the workspace."
                )

            content_length = response.headers.get("Content-Length")
            if content_length and int(content_length) > max_bytes:
                raise ValueError(f"Slack file exceeds configured download limit ({max_bytes} bytes).")

            total = 0
            with destination.open("wb") as output:
                while True:
                    chunk = response.read(1024 * 1024)
                    if not chunk:
                        break
                    total += len(chunk)
                    if total > max_bytes:
                        raise ValueError(
                            f"Slack file exceeds configured download limit ({max_bytes} bytes)."
                        )
                    output.write(chunk)
            _validate_downloaded_content(destination, expected_type)
    except Exception:
        destination.unlink(missing_ok=True)
        raise


def _validate_downloaded_content(path: Path, expected_mimetype: str) -> None:
    """Reject auth HTML and obvious type confusion before OCR or document decoding."""
    expected = (expected_mimetype or "").lower()
    suffix = path.suffix.lower()
    head = path.read_bytes()[:512]
    stripped = head.lstrip().lower()
    expected_is_html = expected in {"text/html", "application/xhtml+xml"} or suffix in {".htm", ".html"}
    if not expected_is_html and (
        stripped.startswith(b"<!doctype html")
        or stripped.startswith(b"<html")
        or (b"<title" in stripped and (b"slack" in stripped or b"sign in" in stripped))
    ):
        raise ValueError(
            "Slack returned an HTML sign-in page instead of the requested file. "
            "The installed Slack app must include files:read and be reinstalled to the workspace."
        )

    is_image = expected.startswith("image/") or suffix in IMAGE_EXTENSIONS
    if is_image and not _matches_image_signature(head, suffix, expected):
        raise ValueError("Downloaded Slack image content does not match a supported image signature.")
    if (expected == "application/pdf" or suffix == ".pdf") and not head.startswith(b"%PDF-"):
        raise ValueError("Downloaded Slack PDF content does not have a valid PDF signature.")
    if suffix in DOCX_EXTENSIONS | PPTX_EXTENSIONS | SPREADSHEET_EXTENSIONS | ZIP_EXTENSIONS and not head.startswith(b"PK"):
        raise ValueError("Downloaded Slack archive/document does not have a valid ZIP signature.")
    if (
        suffix in LEGACY_SPREADSHEET_EXTENSIONS
        or expected == "application/vnd.ms-excel"
    ) and not head.startswith(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"):
        raise ValueError("Downloaded legacy Excel content does not have a valid OLE signature.")


def _matches_image_signature(head: bytes, suffix: str, expected: str) -> bool:
    if head.startswith(b"\xff\xd8\xff"):
        return suffix in {".jpg", ".jpeg", ""} or expected in {"", "image/jpeg"}
    if head.startswith(b"\x89PNG\r\n\x1a\n"):
        return suffix in {".png", ""} or expected in {"", "image/png"}
    if head.startswith((b"GIF87a", b"GIF89a")):
        return suffix in {".gif", ""} or expected in {"", "image/gif"}
    if head.startswith(b"BM"):
        return suffix in {".bmp", ""} or expected in {"", "image/bmp"}
    if head.startswith((b"II*\x00", b"MM\x00*")):
        return suffix in {".tif", ".tiff", ""} or expected in {"", "image/tiff"}
    if len(head) >= 12 and head[:4] == b"RIFF" and head[8:12] == b"WEBP":
        return suffix in {".webp", ""} or expected in {"", "image/webp"}
    if len(head) >= 12 and head[4:8] == b"ftyp" and head[8:12] in {
        b"heic", b"heix", b"hevc", b"hevx", b"mif1", b"msf1",
    }:
        return suffix in {".heic", ".heif", ""} or expected in {"", "image/heic", "image/heif"}
    return False


def _event_files(event: dict[str, Any]) -> list[dict[str, Any]]:
    raw_files: list[dict[str, Any]] = []
    files = event.get("files") or []
    if isinstance(files, list):
        raw_files.extend(item for item in files if isinstance(item, dict))
    single = event.get("file")
    if isinstance(single, dict):
        raw_files.append(single)

    seen: set[str] = set()
    deduped: list[dict[str, Any]] = []
    for item in raw_files:
        key = str(item.get("id") or item.get("url_private") or item.get("name") or id(item))
        if key in seen:
            continue
        seen.add(key)
        deduped.append(item)
    return deduped


def _hydrate_file_meta(file_meta: dict[str, Any], client: Any) -> dict[str, Any]:
    if file_meta.get("url_private") or file_meta.get("url_private_download"):
        return file_meta
    file_id = file_meta.get("id")
    if not file_id or not hasattr(client, "files_info"):
        return file_meta
    try:
        response = client.files_info(file=file_id)
        hydrated = response.get("file") or {}
    except Exception:
        return file_meta
    merged = dict(file_meta)
    for key, value in hydrated.items():
        if not merged.get(key):
            merged[key] = value
    return merged


def _file_record(file_meta: dict[str, Any]) -> dict[str, Any]:
    return {
        "id": file_meta.get("id") or "",
        "name": file_meta.get("name") or file_meta.get("title") or "",
        "title": file_meta.get("title") or "",
        "mimetype": file_meta.get("mimetype") or file_meta.get("filetype") or "",
        "size": file_meta.get("size") or "",
        "url_private": file_meta.get("url_private") or "",
        "url_private_download": file_meta.get("url_private_download") or "",
    }


def _config_int(defaults: dict[str, Any], key: str, default: int) -> int:
    try:
        return max(0, int(defaults.get(key, default)))
    except (TypeError, ValueError):
        return default


def _config_bool(defaults: dict[str, Any], key: str, default: bool) -> bool:
    value = defaults.get(key, default)
    if isinstance(value, bool):
        return value
    if value is None:
        return default
    return str(value).strip().lower() in {"1", "true", "yes", "on"}
