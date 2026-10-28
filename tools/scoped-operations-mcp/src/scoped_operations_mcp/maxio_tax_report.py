from __future__ import annotations

import calendar
import hashlib
import os
import re
import tempfile
from datetime import date, datetime, timedelta
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from pathlib import Path
from typing import Any, Mapping

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.worksheet.page import PageMargins


REPORT_TITLE = "NOVA Sales Tax Report"
REPORT_HEADERS = [
    "Invoice Date",
    "Invoice Number",
    "Customer Name",
    "Address",
    "City",
    "State",
    "ZIP",
    "Sales Amount",
    "Tax Amount",
]
MAX_REPORT_DAYS = 366
CENT = Decimal("0.01")

# These states were independently verified in Maxio Tax History when the
# invoice export snapshot did not carry its billing-state value. They are data
# corrections only and never appear as narrative text in the workbook.
VERIFIED_BILLING_STATE_OVERRIDES = {
    "INV-260762": "MD",
    "INV-260772": "WA",
    "INV-260773": "WA",
    "INV-260780": "TX",
    "INV-260791": "PA",
    "INV-260813": "AZ",
    "INV-260874": "TX",
    "INV-260888": "NY",
    "INV-260931": "TX",
    "INV-261008": "TX",
    "INV-261149": "TX",
    "INV-261168": "TX",
}

STATE_NAMES = {
    "alabama": "AL", "alaska": "AK", "arizona": "AZ", "arkansas": "AR",
    "california": "CA", "colorado": "CO", "connecticut": "CT", "delaware": "DE",
    "district of columbia": "DC", "florida": "FL", "georgia": "GA", "hawaii": "HI",
    "idaho": "ID", "illinois": "IL", "indiana": "IN", "iowa": "IA", "kansas": "KS",
    "kentucky": "KY", "louisiana": "LA", "maine": "ME", "maryland": "MD",
    "massachusetts": "MA", "michigan": "MI", "minnesota": "MN", "mississippi": "MS",
    "missouri": "MO", "montana": "MT", "nebraska": "NE", "nevada": "NV",
    "new hampshire": "NH", "new jersey": "NJ", "new mexico": "NM", "new york": "NY",
    "north carolina": "NC", "north dakota": "ND", "ohio": "OH", "oklahoma": "OK",
    "oregon": "OR", "pennsylvania": "PA", "rhode island": "RI",
    "south carolina": "SC", "south dakota": "SD", "tennessee": "TN", "texas": "TX",
    "utah": "UT", "vermont": "VT", "virginia": "VA", "washington": "WA",
    "west virginia": "WV", "wisconsin": "WI", "wyoming": "WY",
}


class MaxioTaxReportError(ValueError):
    """The requested period or source data cannot produce a reliable report."""


def resolve_report_period(
    *,
    month: str = "",
    start_date: str = "",
    end_date: str = "",
    today: date | None = None,
) -> tuple[date, date]:
    clean_month = str(month or "").strip()
    start_text = str(start_date or "").strip()
    end_text = str(end_date or "").strip()
    if clean_month and (start_text or end_text):
        raise MaxioTaxReportError("Use month or start_date/end_date, not both")
    if bool(start_text) != bool(end_text):
        raise MaxioTaxReportError("start_date and end_date must be provided together")

    if clean_month:
        if not re.fullmatch(r"\d{4}-\d{2}", clean_month):
            raise MaxioTaxReportError("month must use YYYY-MM format")
        try:
            start = datetime.strptime(clean_month, "%Y-%m").date().replace(day=1)
        except ValueError as exc:
            raise MaxioTaxReportError("month must be a valid calendar month") from exc
        end = start.replace(day=calendar.monthrange(start.year, start.month)[1])
    elif start_text:
        try:
            start = datetime.strptime(start_text, "%Y-%m-%d").date()
            end = datetime.strptime(end_text, "%Y-%m-%d").date()
        except ValueError as exc:
            raise MaxioTaxReportError("start_date and end_date must use YYYY-MM-DD format") from exc
    else:
        reference = today or date.today()
        end = reference.replace(day=1) - timedelta(days=1)
        start = end.replace(day=1)

    if end < start:
        raise MaxioTaxReportError("end_date cannot be before start_date")
    if (end - start).days + 1 > MAX_REPORT_DAYS:
        raise MaxioTaxReportError(f"report period cannot exceed {MAX_REPORT_DAYS} days")
    return start, end


def transform_invoices(
    invoices: list[dict[str, Any]],
    start: date,
    end: date,
    *,
    state_overrides: Mapping[str, str] | None = None,
) -> list[dict[str, Any]]:
    overrides = state_overrides if state_overrides is not None else VERIFIED_BILLING_STATE_OVERRIDES
    rows: list[dict[str, Any]] = []
    missing_states: list[str] = []

    for invoice in invoices:
        issue_date = _parse_iso_date(invoice.get("issue_date"))
        tax = _cents(_money(invoice.get("tax_amount")))
        if issue_date < start or issue_date > end or tax <= 0:
            continue
        if str(invoice.get("status") or "").strip().casefold() != "paid":
            continue
        consolidation_level = str(
            invoice.get("consolidation_level") or ""
        ).strip().casefold()
        if (
            consolidation_level == "child"
            or invoice.get("parent_invoice_number")
            or invoice.get("parent_invoice_uid")
        ):
            continue

        number = str(invoice.get("number") or invoice.get("uid") or "").strip()
        subtotal = _money(invoice.get("subtotal_amount"))
        discount = _money(invoice.get("discount_amount"))
        sales = _cents(subtotal - discount)
        total = _cents(_money(invoice.get("total_amount")))
        if total != _cents(sales + tax):
            raise MaxioTaxReportError(
                f"{number or '<blank invoice>'}: invoice total does not equal sales plus tax"
            )

        billing_raw = invoice.get("billing_address") or {}
        billing = billing_raw if isinstance(billing_raw, Mapping) else {}
        customer_raw = invoice.get("customer") or {}
        customer = customer_raw if isinstance(customer_raw, Mapping) else {}
        state_override = overrides.get(number)
        if not state_override and number and not number.upper().startswith("INV-"):
            state_override = overrides.get(f"INV-{number}")
        state = normalize_state(billing.get("state")) or normalize_state(state_override)
        if not state:
            missing_states.append(number or "<blank invoice>")
            continue

        rows.append(
            {
                "invoice_date": issue_date,
                "invoice_number": number,
                "customer_name": _customer_name(customer, billing),
                "address": _address_line(billing),
                "city": str(billing.get("city") or "").strip(),
                "state": state,
                "zip": str(billing.get("zip") or billing.get("postal_code") or "").strip(),
                "sales_amount": sales,
                "tax_amount": tax,
            }
        )

    if missing_states:
        raise MaxioTaxReportError(
            "Taxable invoices have blank billing states: " + ", ".join(sorted(missing_states))
        )
    return _validate_rows(rows)


def write_tax_report(
    rows: list[dict[str, Any]],
    start: date,
    end: date,
    output_dir: Path,
) -> dict[str, Any]:
    output_root = Path(output_dir).resolve()
    output_root.mkdir(parents=True, exist_ok=True, mode=0o770)
    filename = _download_name(start, end)
    destination = output_root / filename
    workbook = _build_workbook(rows, start, end)

    descriptor, temporary_name = tempfile.mkstemp(
        prefix=f".{destination.stem}.",
        suffix=".tmp.xlsx",
        dir=output_root,
    )
    os.close(descriptor)
    temporary = Path(temporary_name)
    try:
        workbook.save(temporary)
        _validate_workbook(temporary, rows, start, end)
        os.replace(temporary, destination)
        os.chmod(destination, 0o660)
        _validate_workbook(destination, rows, start, end)
    finally:
        workbook.close()
        if temporary.exists():
            temporary.unlink()

    content = destination.read_bytes()
    return {
        "ok": True,
        "operation": "tax_report",
        "artifactPath": str(destination),
        "downloadFilename": filename,
        "sha256": hashlib.sha256(content).hexdigest(),
        "sizeBytes": len(content),
        "periodStart": start.isoformat(),
        "periodEnd": end.isoformat(),
        "invoiceCount": len(rows),
        "salesAmount": f"{sum((row['sales_amount'] for row in rows), Decimal('0')):.2f}",
        "taxAmount": f"{sum((row['tax_amount'] for row in rows), Decimal('0')):.2f}",
        "providerWritesPerformed": False,
        "shippingAddressUsed": False,
    }


def normalize_state(value: Any) -> str:
    text = str(value or "").strip()
    if not text:
        return ""
    return STATE_NAMES.get(text.casefold(), text.upper())


def _money(value: Any) -> Decimal:
    try:
        cleaned = str(value if value is not None else "0").strip().replace("$", "").replace(",", "")
        return Decimal(cleaned or "0")
    except (InvalidOperation, ValueError) as exc:
        raise MaxioTaxReportError(f"Invalid monetary value: {value!r}") from exc


def _cents(value: Decimal) -> Decimal:
    return value.quantize(CENT, rounding=ROUND_HALF_UP)


def _parse_iso_date(value: Any) -> date:
    text = str(value or "").strip()
    try:
        return datetime.strptime(text[:10], "%Y-%m-%d").date()
    except ValueError as exc:
        raise MaxioTaxReportError(f"Unsupported invoice date: {text or '<blank>'}") from exc


def _customer_name(customer: Mapping[str, Any], billing: Mapping[str, Any]) -> str:
    organization = str(customer.get("organization") or "").strip()
    if organization:
        return organization
    customer_person = " ".join(
        part for part in (
            str(customer.get("first_name") or "").strip(),
            str(customer.get("last_name") or "").strip(),
        ) if part
    )
    if customer_person:
        return customer_person
    return " ".join(
        part for part in (
            str(billing.get("first_name") or "").strip(),
            str(billing.get("last_name") or "").strip(),
        ) if part
    )


def _address_line(billing: Mapping[str, Any]) -> str:
    return ", ".join(
        part for part in (
            str(billing.get("street") or billing.get("address") or "").strip(),
            str(billing.get("line2") or billing.get("address_2") or "").strip(),
        ) if part
    )


def _validate_rows(rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    if not rows:
        raise MaxioTaxReportError("No invoices with tax greater than zero were found")
    invoice_numbers = [row["invoice_number"] for row in rows]
    if any(not value for value in invoice_numbers):
        raise MaxioTaxReportError("An invoice number is blank")
    if len(invoice_numbers) != len(set(invoice_numbers)):
        raise MaxioTaxReportError("Duplicate invoice numbers were found")
    required = ("customer_name", "address", "city", "state", "zip")
    for row in rows:
        missing = [field for field in required if not row.get(field)]
        if missing:
            raise MaxioTaxReportError(
                f"{row['invoice_number']}: blank required billing fields: {', '.join(missing)}"
            )
    rows.sort(key=lambda item: (item["invoice_date"], item["invoice_number"]))
    return rows


def _period_label(start: date, end: date) -> str:
    return f"{start.strftime('%B')} {start.day}, {start.year} - {end.strftime('%B')} {end.day}, {end.year}"


def _download_name(start: date, end: date) -> str:
    is_month = (
        start.day == 1
        and start.year == end.year
        and start.month == end.month
        and end.day == calendar.monthrange(start.year, start.month)[1]
    )
    if is_month:
        return f"Nova_Monthly_Tax_Report_{start:%Y-%m}.xlsx"
    return f"Nova_Tax_Report_{start:%Y-%m-%d}_to_{end:%Y-%m-%d}.xlsx"


def _build_workbook(rows: list[dict[str, Any]], start: date, end: date) -> Workbook:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Tax Report"

    sheet.merge_cells("A1:I1")
    sheet["A1"] = REPORT_TITLE
    sheet["A1"].font = Font(name="Arial", size=14, bold=True, color="000000")
    sheet["A1"].alignment = Alignment(horizontal="left", vertical="center")

    sheet.merge_cells("A2:I2")
    sheet["A2"] = _period_label(start, end)
    sheet["A2"].font = Font(name="Arial", size=10, color="000000")
    sheet["A2"].alignment = Alignment(horizontal="left", vertical="center")

    header_border = Border(bottom=Side(style="thin", color="000000"))
    for column, label in enumerate(REPORT_HEADERS, start=1):
        cell = sheet.cell(row=4, column=column, value=label)
        cell.font = Font(name="Arial", size=10, bold=True, color="000000")
        cell.alignment = Alignment(horizontal="left", vertical="center")
        cell.border = header_border

    for row_index, item in enumerate(rows, start=5):
        values = [
            item["invoice_date"], item["invoice_number"], item["customer_name"], item["address"],
            item["city"], item["state"], item["zip"], float(item["sales_amount"]), float(item["tax_amount"]),
        ]
        for column, value in enumerate(values, start=1):
            cell = sheet.cell(row=row_index, column=column, value=value)
            cell.font = Font(name="Arial", size=10, color="000000")
            cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=column in {3, 4})
        sheet.cell(row=row_index, column=1).number_format = "mm/dd/yyyy"
        sheet.cell(row=row_index, column=2).number_format = "@"
        sheet.cell(row=row_index, column=6).number_format = "@"
        sheet.cell(row=row_index, column=7).number_format = "@"
        for money_column in (8, 9):
            sheet.cell(row=row_index, column=money_column).number_format = "$#,##0.00"
            sheet.cell(row=row_index, column=money_column).alignment = Alignment(horizontal="right", vertical="top")

    for column, width in {
        "A": 13, "B": 18, "C": 30, "D": 38, "E": 20, "F": 9, "G": 12, "H": 16, "I": 14,
    }.items():
        sheet.column_dimensions[column].width = width
    sheet.row_dimensions[1].height = 22
    sheet.row_dimensions[2].height = 18
    sheet.row_dimensions[4].height = 20
    for row_index in range(5, 5 + len(rows)):
        sheet.row_dimensions[row_index].height = 20

    sheet.freeze_panes = "A5"
    sheet.auto_filter.ref = f"A4:I{4 + len(rows)}"
    sheet.sheet_view.showGridLines = True
    sheet.print_title_rows = "4:4"
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_LETTER
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_margins = PageMargins(left=0.35, right=0.35, top=0.5, bottom=0.5, header=0.2, footer=0.2)
    sheet.print_area = f"A1:I{4 + len(rows)}"
    workbook.properties.title = REPORT_TITLE
    workbook.properties.subject = _period_label(start, end)
    return workbook


def _validate_workbook(path: Path, rows: list[dict[str, Any]], start: date, end: date) -> None:
    workbook = load_workbook(path, data_only=False, read_only=False)
    try:
        if workbook.sheetnames != ["Tax Report"]:
            raise MaxioTaxReportError("Unexpected workbook sheets")
        sheet = workbook["Tax Report"]
        if sheet["A1"].value != REPORT_TITLE or sheet["A2"].value != _period_label(start, end):
            raise MaxioTaxReportError("Workbook title or dates did not match")
        headers = [sheet.cell(row=4, column=column).value for column in range(1, 10)]
        if headers != REPORT_HEADERS:
            raise MaxioTaxReportError("Workbook headers did not match")
        if sheet.max_row != len(rows) + 4 or sheet.max_column != 9:
            raise MaxioTaxReportError("Workbook dimensions did not match")
        invoice_numbers = [str(sheet.cell(row=index, column=2).value) for index in range(5, sheet.max_row + 1)]
        if invoice_numbers != [row["invoice_number"] for row in rows]:
            raise MaxioTaxReportError("Workbook invoice order did not match")
        for cells in sheet.iter_rows():
            for cell in cells:
                if isinstance(cell.value, str) and cell.value.startswith("="):
                    raise MaxioTaxReportError(f"Unexpected formula in {cell.coordinate}")
        sales_total = _cents(sum(
            (Decimal(str(sheet.cell(row=index, column=8).value)) for index in range(5, sheet.max_row + 1)),
            Decimal("0"),
        ))
        tax_total = _cents(sum(
            (Decimal(str(sheet.cell(row=index, column=9).value)) for index in range(5, sheet.max_row + 1)),
            Decimal("0"),
        ))
        expected_sales = _cents(sum((row["sales_amount"] for row in rows), Decimal("0")))
        expected_tax = _cents(sum((row["tax_amount"] for row in rows), Decimal("0")))
        if sales_total != expected_sales or tax_total != expected_tax:
            raise MaxioTaxReportError("Workbook monetary totals did not match")
    finally:
        workbook.close()
