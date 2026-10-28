from __future__ import annotations

import json
from datetime import date
from pathlib import Path

import pytest
from openpyxl import load_workbook

from scoped_operations_mcp.maxio_client import MaxioSettings, ScopedMaxioClient
from scoped_operations_mcp.maxio_server import build_mcp
from scoped_operations_mcp.maxio_tax_report import (
    MaxioTaxReportError,
    REPORT_HEADERS,
    REPORT_TITLE,
    resolve_report_period,
    transform_invoices,
)


def _invoice(number: str = "INV-1", state: str = "MD") -> dict:
    return {
        "issue_date": "2026-07-25",
        "number": number,
        "status": "paid",
        "consolidation_level": "none",
        "parent_invoice_number": None,
        "parent_invoice_uid": None,
        "customer": {
            "organization": "Nova Test Shop",
            "first_name": "Billing",
            "last_name": "Contact",
        },
        "billing_address": {
            "street": "100 Main Street",
            "line2": "Suite 2",
            "city": "Baltimore",
            "state": state,
            "zip": "21201",
        },
        "shipping_address": {
            "street": "999 Shipping Lane",
            "city": "Seattle",
            "state": "WA",
            "zip": "98101",
        },
        "subtotal_amount": "100.00",
        "discount_amount": "10.00",
        "tax_amount": "8.19",
        "total_amount": "98.19",
    }


class FakeResponse:
    def __init__(self, payload) -> None:
        self._payload = payload
        self.content = json.dumps(payload).encode("utf-8")

    def raise_for_status(self) -> None:
        return None

    def json(self):
        return self._payload


class TaxReportSession:
    def __init__(self, invoices: list[dict]) -> None:
        self.invoices = invoices
        self.calls = []

    def get(self, url, **kwargs):
        self.calls.append((url, kwargs))
        assert url.endswith("/invoices.json")
        return FakeResponse(
            {
                "invoices": [dict(row) for row in self.invoices],
                "meta": {"current_page": 1, "total_pages": 1},
            }
        )


def test_period_supports_prior_month_specific_month_and_exact_range() -> None:
    assert resolve_report_period(today=date(2026, 8, 4)) == (
        date(2026, 7, 1),
        date(2026, 7, 31),
    )
    assert resolve_report_period(month="2026-07") == (
        date(2026, 7, 1),
        date(2026, 7, 31),
    )
    assert resolve_report_period(start_date="2026-07-21", end_date="2026-08-03") == (
        date(2026, 7, 21),
        date(2026, 8, 3),
    )


@pytest.mark.parametrize(
    "arguments",
    [
        {"month": "2026-07", "start_date": "2026-07-01", "end_date": "2026-07-31"},
        {"month": "2026-13"},
        {"start_date": "2026-07-01"},
        {"start_date": "2026-08-01", "end_date": "2026-07-01"},
    ],
)
def test_period_rejects_invalid_input(arguments: dict) -> None:
    with pytest.raises(MaxioTaxReportError):
        resolve_report_period(**arguments)


def test_transform_uses_billing_not_shipping_and_requires_reconciled_totals() -> None:
    rows = transform_invoices([_invoice()], date(2026, 7, 21), date(2026, 8, 3))
    assert rows[0]["address"] == "100 Main Street, Suite 2"
    assert rows[0]["city"] == "Baltimore"
    assert rows[0]["state"] == "MD"
    assert str(rows[0]["sales_amount"]) == "90.00"
    assert str(rows[0]["tax_amount"]) == "8.19"

    invalid = _invoice()
    invalid["total_amount"] = "99.00"
    with pytest.raises(MaxioTaxReportError, match="sales plus tax"):
        transform_invoices([invalid], date(2026, 7, 21), date(2026, 8, 3))


def test_transform_includes_only_paid_unconsolidated_invoices() -> None:
    paid = _invoice("INV-PAID")
    voided = _invoice("INV-VOIDED")
    voided["status"] = "voided"
    child = _invoice("INV-CHILD")
    child["consolidation_level"] = "child"
    child["parent_invoice_number"] = "INV-PARENT"

    rows = transform_invoices(
        [paid, voided, child],
        date(2026, 7, 21),
        date(2026, 8, 3),
    )

    assert [row["invoice_number"] for row in rows] == ["INV-PAID"]


def test_verified_state_override_is_data_only_and_unknown_blank_state_stops() -> None:
    verified = _invoice("INV-260813", state="")
    rows = transform_invoices([verified], date(2026, 7, 21), date(2026, 8, 3))
    assert rows[0]["state"] == "AZ"

    live_number = _invoice("260813", state="")
    live_rows = transform_invoices([live_number], date(2026, 7, 21), date(2026, 8, 3))
    assert live_rows[0]["state"] == "AZ"

    with pytest.raises(MaxioTaxReportError, match="blank billing states"):
        transform_invoices(
            [_invoice("INV-UNKNOWN", state="")],
            date(2026, 7, 21),
            date(2026, 8, 3),
        )


def test_typed_broker_creates_plain_validated_xlsx(tmp_path: Path) -> None:
    session = TaxReportSession([_invoice()])
    client = ScopedMaxioClient(
        MaxioSettings(
            "https://example.chargify.com",
            "secret",
            journal_path=tmp_path / "operations.jsonl",
            report_output_dir=tmp_path / "reports",
        ),
        session=session,
    )

    result = client.tax_report(start_date="2026-07-21", end_date="2026-08-03")

    assert result["ok"] is True
    assert result["providerWritesPerformed"] is False
    assert result["shippingAddressUsed"] is False
    assert result["invoiceCount"] == 1
    assert result["salesAmount"] == "90.00"
    assert result["taxAmount"] == "8.19"
    assert len(result["sha256"]) == 64
    assert session.calls[0][1]["params"]["date_field"] == "issue_date"
    assert session.calls[0][1]["params"]["start_date"] == "2026-07-21"
    assert session.calls[0][1]["params"]["end_date"] == "2026-08-03"

    report_path = Path(result["artifactPath"])
    assert report_path.name == "Nova_Tax_Report_2026-07-21_to_2026-08-03.xlsx"
    workbook = load_workbook(report_path, data_only=False)
    try:
        assert workbook.sheetnames == ["Tax Report"]
        sheet = workbook["Tax Report"]
        assert sheet["A1"].value == REPORT_TITLE
        assert sheet["A2"].value == "July 21, 2026 - August 3, 2026"
        assert [sheet.cell(row=4, column=column).value for column in range(1, 10)] == REPORT_HEADERS
        assert sheet.max_row == 5
        assert sheet.max_column == 9
        assert sheet["B5"].value == "INV-1"
        assert sheet["D5"].value == "100 Main Street, Suite 2"
        text_values = "\n".join(
            str(cell.value)
            for row in sheet.iter_rows()
            for cell in row
            if cell.value is not None
        )
        assert "Report scope" not in text_values
        assert "Confidential" not in text_values
        assert "Maxio Tax History" not in text_values
        assert all(
            not (isinstance(cell.value, str) and cell.value.startswith("="))
            for row in sheet.iter_rows()
            for cell in row
        )
    finally:
        workbook.close()


def test_tax_report_tool_is_available_to_billing_and_coding_lanes(tmp_path: Path) -> None:
    client = ScopedMaxioClient(
        MaxioSettings(
            "https://example.chargify.com",
            "secret",
            report_output_dir=tmp_path,
        ),
        session=TaxReportSession([_invoice()]),
    )
    for lane in ("billing_agent", "coding_agent"):
        mcp = build_mcp(lane, client)
        assert "maxio_tax_report" in mcp._tool_manager._tools
        capabilities = mcp._tool_manager._tools["maxio_capabilities"].fn()
        assert "maxio_tax_report" in capabilities["operations"]
